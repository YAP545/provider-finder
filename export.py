"""Step 4: write CSV + formatted Excel (All + Shortlist sheets)."""
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

COLUMNS = [
    "rank", "shortlisted", "suitability_score", "name", "name_raw", "website", "category", "provider_type", "country",
    "country_source", "services", "key_skills", "portfolio_urls", "pricing_info", "email", "phone",
    "social_links", "summary", "experience_evidence", "red_flags", "skill_relevance", "experience_score",
    "portfolio_score", "transparency_score", "confidence", "ai_used", "scored_by", "source_query", "pages_checked",
    "collected_at",
    "experience_source_url", "experience_status",
    "portfolio_source_url", "portfolio_status",
    "pricing_source_url", "pricing_status",
    "last_checked", "strengths", "limitations",
]


def export(rows: list, out_dir: str = "data") -> pd.DataFrame:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    
    # Ensure every column exists on every row dict
    normalized_rows = []
    for r in rows:
        row_dict = {}
        for c in COLUMNS:
            row_dict[c] = r.get(c, "")
        normalized_rows.append(row_dict)

    df = pd.DataFrame(normalized_rows)[COLUMNS]
    df.to_csv(out / "providers.csv", index=False, encoding="utf-8-sig")
    
    # Ensure boolean shortlisted
    shortlisted_mask = df["shortlisted"].astype(str).str.lower().isin(["true", "1"])
    shortlist_df = df[shortlisted_mask]

    with pd.ExcelWriter(out / "providers.xlsx", engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="All providers", index=False)
        shortlist_df.to_excel(xw, sheet_name="Shortlist", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(name="Arial", bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F3864")
            ws.freeze_panes = "E2"
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value or "")) for c in col[:50])
                ws.column_dimensions[get_column_letter(i)].width = min(45, max(10, width + 2))
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.font = Font(name="Arial", size=10)
                    c.alignment = Alignment(vertical="top", wrap_text=False)
    return df
